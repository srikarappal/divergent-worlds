from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils.cell import get_column_letter, range_boundaries

from support.guest_fakes import build_agent

from rlworld.guest import GuestFailure
from rlworld.guest.excel import ExcelIntegration, ExcelNative
from rlworld.guest.protocol import NativeCallError
from rlworld.worlds.generator import generate_mid_market_company
from rlworld.worlds.integrations import DEFAULT_INTEGRATIONS, UnsupportedAppError


class FakeComError(Exception):
    """What Excel's COM objects raise, such as for a range address it cannot parse."""


class FakeRange:
    """A rectangle of a fake sheet with Excel's Value, Formula, Address and Resize."""

    def __init__(self, sheet, min_column, min_row, max_column, max_row):
        self.sheet = sheet
        self.min_column, self.min_row, self.max_column, self.max_row = min_column, min_row, max_column, max_row

    def coordinates(self):
        return [[(row, column) for column in range(self.min_column, self.max_column + 1)]
                for row in range(self.min_row, self.max_row + 1)]

    def grid(self, field, empty):
        rows = tuple(tuple(self.sheet.cells.get(coordinate, {}).get(field, empty) for coordinate in row)
                     for row in self.coordinates())
        return rows[0][0] if len(rows) == 1 and len(rows[0]) == 1 else rows

    @property
    def Value(self):
        return self.grid("value", None)

    @property
    def Formula(self):
        return self.grid("formula", "")

    @Formula.setter
    def Formula(self, values):
        for row, written in zip(self.coordinates(), values, strict=True):
            for coordinate, text in zip(row, written, strict=True):
                computed = self.sheet.calculator.get(text, None) if str(text).startswith("=") else text
                self.sheet.cells[coordinate] = {"formula": text, "value": computed}

    @property
    def Address(self):
        start = f"${get_column_letter(self.min_column)}${self.min_row}"
        if (self.min_column, self.min_row) == (self.max_column, self.max_row):
            return start
        return f"{start}:${get_column_letter(self.max_column)}${self.max_row}"

    def Resize(self, rows, columns):
        return FakeRange(self.sheet, self.min_column, self.min_row, self.min_column + columns - 1,
                         self.min_row + rows - 1)


class FakeSheet:
    def __init__(self, name, values=None, calculator=None):
        self.Name = name
        self.calculator = calculator or {}
        self.cells = {}
        for address, value in (values or {}).items():
            min_column, min_row, _, _ = range_boundaries(address)
            self.cells[(min_row, min_column)] = {"formula": value, "value": value}

    def Range(self, address):
        try:
            return FakeRange(self, *range_boundaries(address))
        except (TypeError, ValueError) as error:
            raise FakeComError(f"Exception occurred: Range of {address} failed") from error


class FakeWorkbook:
    def __init__(self, name, full_name, sheets):
        self.Name = name
        self.FullName = full_name
        self.Worksheets = list(sheets)
        self.ActiveSheet = self.Worksheets[0]
        self.saves = 0
        self.copies = []
        self.closed = None

    def Save(self):
        self.saves += 1

    def SaveCopyAs(self, path):
        """Writes what Excel last computed, the way a saved workbook caches its formulas' values."""
        self.copies.append(path)
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        for sheet in self.Worksheets:
            written = workbook.create_sheet(sheet.Name)
            for (row, column), cell in sheet.cells.items():
                written.cell(row=row, column=column, value=cell["value"])
        workbook.save(path)

    def Close(self, save_changes):
        self.closed = save_changes


class FakeWorkbooks(list):
    @property
    def Count(self):
        return len(self)


class FakeExcel:
    def __init__(self, workbooks):
        self.Workbooks = FakeWorkbooks(workbooks)
        self.ActiveWorkbook = workbooks[0] if workbooks else None


class RunningExcel:
    """Stands in for asking the session for its running Excel."""

    def __init__(self, excel):
        self.excel = excel

    def application(self):
        return self.excel


def budget(home):
    totals = FakeSheet("Totals", {"A1": "open", "B1": 1250.5, "A2": "overdue", "B2": 310},
                       calculator={"=B1+B2": 1560.5})
    notes = FakeSheet("Notes", {"A1": "checked by finance"})
    return FakeWorkbook("budget.xlsx", str(home / "Documents" / "budget.xlsx"), [totals, notes])


@pytest.fixture
def excel(tmp_path):
    return FakeExcel([budget(tmp_path / "home"), FakeWorkbook("Book1", "Book1", [FakeSheet("Sheet1")])])


def native(excel):
    return ExcelNative(RunningExcel(excel).application, (FakeComError,))


def test_native_calls_list_sheets_read_ranges_and_save(excel):
    api = native(excel)
    assert api.call("list_sheets", {}) == {"workbook": "budget.xlsx", "sheets": ["Totals", "Notes"]}
    assert api.call("list_sheets", {"workbook": "Book1"}) == {"workbook": "Book1", "sheets": ["Sheet1"]}
    assert api.call("read_range", {"range": "A1:B2"}) == {
        "sheet": "Totals", "range": "$A$1:$B$2", "values": [["open", 1250.5], ["overdue", 310]],
        "formulas": [["open", 1250.5], ["overdue", 310]]}
    assert api.call("read_range", {"sheet": "Notes", "range": "A1"})["values"] == [["checked by finance"]]
    assert api.call("save", {}) == {"workbook": "budget.xlsx", "saved": True}
    assert excel.Workbooks[0].saves == 1


def test_writes_fill_as_many_cells_as_the_values_and_keep_formulas(excel):
    api = native(excel)
    written = api.call("write_range", {"range": "C1", "values": [["total", "=B1+B2"], ["checked", True]]})
    assert written == {"sheet": "Totals", "range": "$C$1:$D$2"}
    read = api.call("read_range", {"range": "C1:D2"})
    assert read["formulas"] == [["total", "=B1+B2"], ["checked", True]]
    assert read["values"] == [["total", 1560.5], ["checked", True]]


@pytest.mark.parametrize(("method", "arguments", "reason"), [
    ("delete_sheet", {}, "excel has no method delete_sheet"),
    ("read_range", {}, "needs a range"),
    ("read_range", {"range": "not a range"}, "excel refused read_range"),
    ("read_range", {"sheet": "Summary", "range": "A1"}, "budget.xlsx has no sheet named Summary"),
    ("save", {"workbook": "missing.xlsx"}, "no open workbook is named missing.xlsx"),
    ("write_range", {"range": "A1", "values": [[1, 2], [3]]}, "equally long rows"),
    ("write_range", {"range": "A1", "values": []}, "equally long rows"),
    ("write_range", {"range": "A1", "values": "1250"}, "equally long rows"),
])
def test_refused_calls_say_why(excel, method, arguments, reason):
    with pytest.raises(NativeCallError, match=reason):
        native(excel).call(method, arguments)


def test_calls_without_a_workbook_are_refused():
    for excel in (None, FakeExcel([])):
        with pytest.raises(NativeCallError, match="no open workbook"):
            native(excel).call("list_sheets", {})


def test_extraction_saves_copies_and_reads_the_values_excel_computed(excel, tmp_path):
    home = tmp_path / "home"
    shared = FakeWorkbook("q3.xlsm", "/srv/shared/q3.xlsm", [FakeSheet("Q3", {"A1": 7})])
    legacy = FakeWorkbook("old.xls", str(home / "old.xls"), [FakeSheet("Old", {"A1": 1})])
    excel.Workbooks.extend([shared, legacy])
    native(excel).call("write_range", {"range": "C1", "values": [["=B1+B2"]]})
    state = ExcelIntegration(RunningExcel(excel).application, home).extract(home / "world")
    workbooks = state["workbooks"]
    assert set(workbooks) == {"Documents/budget.xlsx", "Book1", "/srv/shared/q3.xlsm", "old.xls"}
    assert workbooks["Documents/budget.xlsx"] == {"kind": "spreadsheet", "sheet": "Totals", "cells": {
        "A1": {"value": "open"}, "B1": {"value": 1250.5}, "A2": {"value": "overdue"}, "B2": {"value": 310},
        "C1": {"value": 1560.5}}}
    assert workbooks["Book1"] == {"kind": "spreadsheet", "sheet": "Sheet1", "cells": {}}
    assert workbooks["/srv/shared/q3.xlsm"]["cells"] == {"A1": {"value": 7}}
    assert [Path(copy).suffix for copy in legacy.copies] == [".xlsx"]
    assert all(workbook.saves == 0 and workbook.closed is None for workbook in excel.Workbooks)
    assert all(not Path(copy).exists() for workbook in excel.Workbooks for copy in workbook.copies)


def test_unchanged_workbooks_extract_the_same_every_time(excel, tmp_path):
    integration = ExcelIntegration(RunningExcel(excel).application, tmp_path / "home")
    assert integration.extract(tmp_path) == integration.extract(tmp_path)


def test_extraction_without_excel_and_resets_close_without_saving(excel, tmp_path):
    assert ExcelIntegration(RunningExcel(None).application, tmp_path).extract(tmp_path) == {"workbooks": {}}
    integration = ExcelIntegration(RunningExcel(excel).application, tmp_path)
    integration.materialize(None, tmp_path)
    integration.reset(tmp_path)
    assert [workbook.closed for workbook in excel.Workbooks] == [False, False]
    ExcelIntegration(RunningExcel(None).application, tmp_path).reset(tmp_path)


def test_the_guest_agent_reports_native_results_in_the_next_observation(excel, tmp_path):
    agent = build_agent(tmp_path, native={"excel": native(excel)})
    agent.handle("act", {"action": {"name": "native_call", "app": "excel", "method": "write_range",
                                    "arguments": {"range": "E5", "values": [[42]]}}})
    observation = agent.handle("observe", {"kinds": []})
    assert observation["native_result"] == {"result": {"sheet": "Totals", "range": "$E$5"}}
    assert "native_result" not in agent.handle("observe", {"kinds": []})
    agent.handle("act", {"action": {"name": "native_call", "app": "excel", "method": "read_range",
                                    "arguments": {"range": "1:1:1"}}})
    refusal = agent.handle("observe", {"kinds": []})["native_result"]
    assert set(refusal) == {"error"} and "excel refused read_range" in refusal["error"]
    with pytest.raises(GuestFailure) as failure:
        agent.handle("act", {"action": {"name": "native_call", "app": "word", "method": "open"}})
    assert failure.value.code == "unsupported_action" and "word" in str(failure.value)


def test_the_guest_agent_materializes_and_extracts_excel_beside_the_world_apps(excel, tmp_path):
    integration = ExcelIntegration(RunningExcel(excel).application, tmp_path / "home")
    agent = build_agent(tmp_path, integrations={**DEFAULT_INTEGRATIONS, "excel": integration}, world_inside_home=True)
    world = generate_mid_market_company(seed=7, size={"employees": 3, "customers": 2, "emails_per_person": 1})
    agent.handle("inject", {"world": world.model_dump(mode="json"), "apps": ["files", "excel"]})
    state = agent.handle("extract", {"apps": ["files", "excel"]})
    assert set(state) == {"files", "excel"} and state["files"]
    assert set(state["excel"]["workbooks"]) == {"Documents/budget.xlsx", "Book1"}
    with pytest.raises(UnsupportedAppError):
        build_agent(tmp_path / "linux").handle("extract", {"apps": ["excel"]})
